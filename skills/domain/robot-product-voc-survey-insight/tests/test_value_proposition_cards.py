from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock
from zipfile import ZipFile


SCRIPT_ROOT = Path(__file__).resolve().parents[1] / "scripts"
if str(SCRIPT_ROOT) in sys.path:
    sys.path.remove(str(SCRIPT_ROOT))
sys.path.insert(0, str(SCRIPT_ROOT))
for module_name in (
    "derive_higher_order_insights",
    "generate_analysis_scaffold",
    "generate_insight_bundle",
    "llm_reasoning_engine",
    "presentation_report_renderer",
    "render_presentation_html",
    "run_multisource_preflight",
    "user_chain_dependency_registry",
):
    sys.modules.pop(module_name, None)

import generate_insight_bundle as insight_bundle


class FakeConnection:
    def close(self) -> None:
        return None


def fake_snapshot(_conn: FakeConnection, spu_id: str) -> dict[str, object]:
    snapshots = {
        "ecovacs_x12": {
            "summary": {"message_count": 120, "negative_count": 18},
            "themes": [{"tag_name": "清洁效果好"}, {"tag_name": "科技外观"}],
            "negatives": [{"tag_name": "拖地水痕"}, {"tag_name": "避障失败"}],
        },
        "ecovacs_x11": {
            "summary": {"message_count": 90, "negative_count": 22},
            "themes": [{"tag_name": "清洁效果好"}],
            "negatives": [{"tag_name": "维护麻烦"}],
        },
        "mova_z70": {
            "summary": {"message_count": 80, "negative_count": 10},
            "themes": [{"tag_name": "避障稳定"}, {"tag_name": "维护省心"}],
            "negatives": [{"tag_name": "噪音偏大"}],
        },
        "ecovacs_t90": {
            "summary": {"message_count": 130, "negative_count": 16},
            "themes": [{"tag_name": "拖地干净"}, {"tag_name": "超薄外观"}],
            "negatives": [{"tag_name": "维护麻烦"}, {"tag_name": "拖地水痕"}],
        },
        "ecovacs_t80s": {
            "summary": {"message_count": 100, "negative_count": 25},
            "themes": [{"tag_name": "拖地干净"}],
            "negatives": [{"tag_name": "滚刷缠绕"}],
        },
        "roborock_g30s": {
            "summary": {"message_count": 70, "negative_count": 9},
            "themes": [{"tag_name": "清洁覆盖完整"}],
            "negatives": [{"tag_name": "基站维护"}],
        },
    }
    return snapshots.get(spu_id, {"summary": {}, "themes": [], "negatives": []})


def build_result(series_code: str) -> dict[str, object]:
    if series_code == "T":
        return {
            "selected_spus": [
                {
                    "spu_id": "ecovacs_t90",
                    "spu_name": "T90 ULTRA",
                    "self_competitor_type": "self",
                    "series_code": "T",
                    "lifecycle_status": "current",
                },
                {
                    "spu_id": "ecovacs_t80s",
                    "spu_name": "T80S",
                    "self_competitor_type": "self",
                    "series_code": "T",
                    "lifecycle_status": "previous",
                },
                {
                    "spu_id": "roborock_g30s",
                    "spu_name": "石头G30s pro",
                    "self_competitor_type": "competitor",
                    "series_code": "T",
                },
            ],
            "voc_rows": [
                {"canonical_spu_id": "ecovacs_t90", "message_count": 130, "negative_count": 16},
                {"canonical_spu_id": "roborock_g30s", "message_count": 70, "negative_count": 9},
            ],
            "survey_rows": [],
            "summary_rows": [],
            "support_matrix": [
                {"module_code": "voc_analysis", "module_name_cn": "VOC分析", "support_level": "支持"},
                {"module_code": "survey", "module_name_cn": "问卷", "support_level": "待补充"},
            ],
        }
    return {
        "selected_spus": [
            {
                "spu_id": "ecovacs_x12",
                "spu_name": "X12 Pro",
                "self_competitor_type": "self",
                "series_code": "X",
                "lifecycle_status": "current",
            },
            {
                "spu_id": "ecovacs_x11",
                "spu_name": "X11 Pro",
                "self_competitor_type": "self",
                "series_code": "X",
                "lifecycle_status": "previous",
            },
            {
                "spu_id": "mova_z70",
                "spu_name": "mova Z70 Pro",
                "self_competitor_type": "competitor",
                "series_code": "X",
            },
        ],
        "voc_rows": [
            {"canonical_spu_id": "ecovacs_x12", "message_count": 120, "negative_count": 18},
            {"canonical_spu_id": "mova_z70", "message_count": 80, "negative_count": 10},
        ],
        "survey_rows": [],
        "summary_rows": [],
        "support_matrix": [
            {"module_code": "voc_analysis", "module_name_cn": "VOC分析", "support_level": "支持"},
            {"module_code": "survey", "module_name_cn": "问卷", "support_level": "待补充"},
        ],
    }


class ValuePropositionCardTests(unittest.TestCase):
    def setUp(self) -> None:
        self.args = SimpleNamespace(pairwise_competitor_spu=None)

    def write_minimal_docx(self, path: Path, paragraphs: list[str]) -> None:
        document_xml = (
            "<w:document><w:body>"
            + "".join(f"<w:p><w:r><w:t>{paragraph}</w:t></w:r></w:p>" for paragraph in paragraphs)
            + "</w:body></w:document>"
        )
        with ZipFile(path, "w") as archive:
            archive.writestr("word/document.xml", document_xml)

    def test_value_translation_reads_dragging_as_result_trust_not_mop_parameter(self) -> None:
        rule = insight_bundle.value_rule_for_text("拖地水痕")
        self.assertIn("清洁结果不稳定", rule["pain_point"])
        self.assertIn("更稳的清洁结果", rule["need"])
        self.assertNotIn("更强拖布", rule["need"])

    def test_x_and_t_default_value_propositions_stay_distinct(self) -> None:
        x_summary = insight_bundle.build_series_value_proposition_summary(
            series_code="X",
            positive_tags=["清洁效果好"],
            negative_tags=["避障失败"],
            competitor_name="mova Z70 Pro",
        )
        t_summary = insight_bundle.build_series_value_proposition_summary(
            series_code="T",
            positive_tags=["拖地干净"],
            negative_tags=["维护麻烦"],
            competitor_name="石头G30s pro",
        )
        self.assertIn("高端场景", x_summary["core_value_proposition"])
        self.assertIn("真托管", x_summary["core_value_proposition"])
        self.assertIn("拖后结果可信", t_summary["buy_not"])
        self.assertIn("主销场景", t_summary["core_value_proposition"])
        self.assertNotEqual(x_summary["core_value_proposition"], t_summary["core_value_proposition"])

    @mock.patch.object(insight_bundle, "fetch_voc_product_snapshot", side_effect=fake_snapshot)
    @mock.patch.object(insight_bundle, "connect", return_value=FakeConnection())
    @mock.patch.object(
        insight_bundle,
        "build_series_competition_engineering_extension",
        return_value={"value_config_test_cards": []},
    )
    @mock.patch.object(
        insight_bundle,
        "build_future_intelligence_extension_artifacts",
        return_value={
            "roadmap_stress_test_cards": {
                "cards": [
                    {
                        "value_axis": "免维护闭环",
                        "awe_industry_endgame_signal": "生态隐形融合",
                        "dreame_2027_2028_pressure": "2028：自动换滚筒 4-6 组",
                        "roadmap_action": "加速",
                        "pm_judgment": "自动换滚筒会把高端托管竞争推向维护闭环。",
                    }
                ]
            },
            "future_competition_strategy_inputs": {
                "x_series_strategy_bias": "X 系优先守住高端托管从容感。",
                "t_series_strategy_bias": "T 系优先守住主销默认答案。",
                "do_not_follow_parameter_wars": ["不把 250℃蒸汽直接当主价值。"],
            },
        },
    )
    def test_render_series_report_contains_product_judgment_sections(
        self,
        _future: mock.Mock,
        _engineering: mock.Mock,
        _connect: mock.Mock,
        _fetch: mock.Mock,
    ) -> None:
        report = insight_bundle.render_series_report(
            build_result("X"),
            analysis_title="X 系列全量本品与竞品用户分析",
            time_scope="2026Q2",
            market="中国",
            analysis_goal=["以产品规划视角读用户"],
        )
        self.assertIn("## 4. 价值主张收口", report)
        self.assertIn("主买点", report)
        self.assertIn("主断点", report)
        self.assertIn("竞品分别在抢哪类用户心智", report)
        self.assertIn("价值主张的竞争证据延伸", report)
        self.assertIn("前瞻情报压力测试", report)
        self.assertIn("对下一代定义的三条启示", report)

    @mock.patch.object(insight_bundle, "fetch_voc_product_snapshot", side_effect=fake_snapshot)
    @mock.patch.object(insight_bundle, "connect", return_value=FakeConnection())
    def test_value_proposition_cards_expose_dashboard_ready_layers(self, _connect: mock.Mock, _fetch: mock.Mock) -> None:
        payload = insight_bundle.build_value_proposition_cards(
            series_payloads=[
                {"series_code": "X", "series_label": "X 系列", "result": build_result("X")},
                {"series_code": "T", "series_label": "T 系列", "result": build_result("T")},
            ],
            args=self.args,
            generated_from="test/main_report.md",
        )
        self.assertEqual(payload["card_count"], 2)
        self.assertEqual(payload["generated_from"]["card_types"], list(insight_bundle.VALUE_PROPOSITION_CARD_TYPES))
        first_card = payload["series_cards"][0]
        self.assertIn("pain_need_translation_cards", first_card)
        self.assertIn("competitor_switch_cards", first_card)
        self.assertIn("next_gen_definition_inputs", first_card)
        self.assertNotIn("超薄", first_card["core_value_proposition"])
        axis_names = [row["value_axis"] for row in first_card["core_value_cards"]]
        self.assertIn("边界项", axis_names)

    def test_artifact_index_knows_value_proposition_cards(self) -> None:
        index = insight_bundle.build_artifact_index(
            [
                Path("/tmp/value_proposition_cards.json"),
                Path("/tmp/integrated_strategy_report.md"),
                Path("/tmp/report_assembly_cards.json"),
                Path("/tmp/competition_engineering_evidence_cards.json"),
                Path("/tmp/value_config_test_map.json"),
                Path("/tmp/competitive_gap_attribution_cards.json"),
                Path("/tmp/next_gen_config_test_inputs.json"),
                Path("/tmp/future_intelligence_signals.json"),
                Path("/tmp/dreame_roadmap_timeline.json"),
                Path("/tmp/roadmap_stress_test_cards.json"),
                Path("/tmp/future_competition_strategy_inputs.json"),
            ]
        )
        self.assertEqual(index["value_proposition_cards"], "/tmp/value_proposition_cards.json")
        self.assertEqual(index["integrated_strategy_report"], "/tmp/integrated_strategy_report.md")
        self.assertEqual(index["report_assembly_cards"], "/tmp/report_assembly_cards.json")
        self.assertEqual(index["competition_engineering_evidence_cards"], "/tmp/competition_engineering_evidence_cards.json")
        self.assertEqual(index["value_config_test_map"], "/tmp/value_config_test_map.json")
        self.assertEqual(index["competitive_gap_attribution_cards"], "/tmp/competitive_gap_attribution_cards.json")
        self.assertEqual(index["next_gen_config_test_inputs"], "/tmp/next_gen_config_test_inputs.json")
        self.assertEqual(index["future_intelligence_signals"], "/tmp/future_intelligence_signals.json")
        self.assertEqual(index["dreame_roadmap_timeline"], "/tmp/dreame_roadmap_timeline.json")
        self.assertEqual(index["roadmap_stress_test_cards"], "/tmp/roadmap_stress_test_cards.json")
        self.assertEqual(index["future_competition_strategy_inputs"], "/tmp/future_competition_strategy_inputs.json")

    def test_config_parser_recognizes_product_columns_and_fields(self) -> None:
        try:
            import openpyxl
        except ModuleNotFoundError:
            self.skipTest("openpyxl is not installed")

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "config.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.cell(row=2, column=3, value="科沃斯")
            ws.cell(row=4, column=3, value="X12")
            ws.cell(row=4, column=4, value="G30S PRO")
            ws.cell(row=16, column=1, value="导航方式")
            ws.cell(row=17, column=2, value="避障类型")
            ws.cell(row=17, column=3, value="线激光")
            ws.cell(row=17, column=4, value="AI视觉")
            wb.save(path)

            payload = insight_bundle.load_product_config_capability_matrix(path)
        self.assertTrue(payload["available"])
        self.assertEqual(payload["product_count"], 2)
        self.assertIn("X12", payload["by_product"])
        self.assertEqual(payload["by_product"]["X12"]["fields"]["导航方式/避障类型"], "线激光")

    def test_test_parser_normalizes_multi_sheet_metric_rows(self) -> None:
        try:
            import openpyxl
        except ModuleNotFoundError:
            self.skipTest("openpyxl is not installed")

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "test.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "X12"
            ws.cell(row=1, column=1, value="测试用例名")
            ws.cell(row=1, column=5, value="X12")
            ws.cell(row=2, column=9, value="数据1")
            ws.cell(row=3, column=1, value="边角覆盖率")
            ws.cell(row=3, column=2, value="黑色踢脚线")
            ws.cell(row=3, column=3, value="标准")
            ws.cell(row=3, column=4, value="覆盖率")
            ws.cell(row=3, column=9, value="98%")
            wb.save(path)

            payload = insight_bundle.load_product_test_metric_matrix(path)
        self.assertTrue(payload["available"])
        self.assertEqual(payload["product_count"], 1)
        self.assertEqual(payload["rows"][0]["test_case"], "边角覆盖率")
        self.assertEqual(payload["rows"][0]["test_value"], "98%")

    def test_competition_engineering_extension_keeps_missing_tests_as_observation(self) -> None:
        config_matrix = {
            "available": True,
            "product_count": 2,
            "by_product": {
                "X12": {
                    "product_name": "X12",
                    "fields": {
                        "扫/真空度(pa)": "11000",
                        "导航方式/避障类型": "线激光",
                    },
                },
                "mova Z70 Pro": {
                    "product_name": "mova Z70 Pro",
                    "fields": {
                        "导航方式/避障类型": "AI视觉",
                    },
                },
            },
        }
        test_matrix = {
            "available": True,
            "product_count": 1,
            "by_product": {
                "X12": [
                    {
                        "test_case": "边角覆盖率",
                        "scene_or_medium": "黑色踢脚线",
                        "mode_or_level": "标准",
                        "data_content": "覆盖率",
                        "test_value": "98%",
                    }
                ]
            },
        }
        payload = insight_bundle.build_series_competition_engineering_extension(
            result=build_result("X"),
            args=self.args,
            series_code="X",
            config_matrix=config_matrix,
            test_matrix=test_matrix,
        )
        self.assertEqual(payload["series_code"], "X")
        self.assertEqual(payload["card_types"], list(insight_bundle.COMPETITION_ENGINEERING_CARD_TYPES))
        cards = payload["value_config_test_cards"]
        self.assertEqual(len(cards), len(insight_bundle.VALUE_CONFIG_TEST_AXIS_RULES))
        competitor_cards = cards[0]["competitor_switch_evidence_cards"]
        self.assertIn("不能写成确定性领先", competitor_cards[0]["switch_reading"])
        self.assertTrue(payload["gap_attribution_cards"])

    def test_value_config_test_map_has_all_five_value_axes(self) -> None:
        payload = insight_bundle.build_value_config_test_map()
        self.assertEqual(payload["axis_count"], 5)
        axis_names = [row["value_axis"] for row in payload["value_axes"]]
        self.assertIn("清洁结果可信", axis_names)
        self.assertIn("免维护闭环", axis_names)
        self.assertIn("真托管少接管", axis_names)

    def test_extract_docx_paragraphs_reads_future_intelligence_docs(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "awe.docx"
            self.write_minimal_docx(path, ["AWE 展示轮足机器人。", "洗扫拖集成站进入平嵌生态。"])

            paragraphs = insight_bundle.extract_docx_paragraphs(path)
        self.assertEqual(paragraphs, ["AWE 展示轮足机器人。", "洗扫拖集成站进入平嵌生态。"])

    def test_future_intelligence_signals_classifies_awe_endgame_themes(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            awe_path = Path(temp_dir) / "awe.docx"
            method_path = Path(temp_dir) / "method.docx"
            self.write_minimal_docx(
                awe_path,
                [
                    "AWE 出现轮足、飞行机器人和多台阶越障，说明空间形态正在越界。",
                    "洗扫拖集成站、全嵌和平嵌正在进入家装生态。",
                    "耗材订阅和云服务订阅开始影响商业模式。",
                ],
            )
            self.write_minimal_docx(method_path, ["AWE 不是展会纪要，而是行业终局信号。"])

            payload = insight_bundle.build_future_intelligence_signals(
                awe_paths=[awe_path],
                method_path=method_path,
            )
        group_names = [row["signal_group"] for row in payload["signal_groups"]]
        self.assertIn("空间形态越界", group_names)
        self.assertIn("生态隐形融合", group_names)
        self.assertIn("商业模式变化", group_names)
        self.assertEqual(payload["red_lines"], list(insight_bundle.FUTURE_INTELLIGENCE_RED_LINES))

    def test_dreame_roadmap_timeline_keeps_unverified_boundary(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "dreame.docx"
            self.write_minimal_docx(
                path,
                [
                    "追觅 2027 规划 7-8cm 超薄、250℃蒸汽、45000Pa 大吸力和多光谱污渍识别。",
                    "追觅 2028 规划自动换滚筒 4-6 组、干湿垃圾自动分离、污水箱自洁。",
                ],
            )

            payload = insight_bundle.build_dreame_roadmap_timeline(path)
        self.assertIn("2027", payload["timeline_by_year"])
        self.assertIn("2028", payload["timeline_by_year"])
        statuses = [row["maturity_status"] for row in payload["rows"]]
        self.assertIn("maturity_unverified", statuses)
        self.assertIn("强竞品意图假设", payload["confidence_boundary"])

    def test_roadmap_stress_test_outputs_all_value_axes_and_actions(self) -> None:
        future_signals = {
            "signal_groups": [
                {"signal_group": "空间形态越界", "value_axes": ["复杂地形通过性"]},
                {"signal_group": "清洁技术前沿", "value_axes": ["清洁结果可信"]},
            ]
        }
        dreame_timeline = {
            "rows": [
                {"year": "2027", "value_axis": "清洁结果可信", "technology_signal": "250℃蒸汽"},
                {"year": "2028", "value_axis": "免维护闭环", "technology_signal": "自动换滚筒 4-6 组"},
            ]
        }
        payload = insight_bundle.build_roadmap_stress_test_cards(future_signals, dreame_timeline)
        axis_names = [row["value_axis"] for row in payload["cards"]]
        actions = [row["roadmap_action"] for row in payload["cards"]]
        self.assertEqual(axis_names, [rule["value_axis"] for rule in insight_bundle.ROADMAP_STRESS_AXIS_RULES])
        self.assertIn("坚持", actions)
        self.assertIn("加速", actions)
        self.assertIn("补洞", actions)
        self.assertIn("观察", actions)

    def test_future_strategy_inputs_keep_parameter_wars_out_of_core_value(self) -> None:
        stress_cards = insight_bundle.build_roadmap_stress_test_cards(
            {"signal_groups": []},
            {"rows": []},
        )
        payload = insight_bundle.build_future_competition_strategy_inputs(stress_cards)
        do_not_follow = " ".join(payload["do_not_follow_parameter_wars"])
        self.assertIn("不把 250℃蒸汽直接当主价值", do_not_follow)
        self.assertIn("X 系优先", payload["x_series_strategy_bias"])
        self.assertIn("T 系优先", payload["t_series_strategy_bias"])

    def test_integrated_report_assembly_keeps_dual_reading_and_boundaries(self) -> None:
        value_proposition_cards = {
            "series_cards": [
                {
                    "series_code": "X",
                    "series_label": "X 系列",
                    "core_value_proposition": "更可信的高端强能力与真托管。",
                    "main_buy_point": "高端场景也能放手。",
                    "main_breakpoint": "避障失败会拉回接管。",
                    "main_switch_point": "竞品抢托管信任。",
                },
                {
                    "series_code": "T",
                    "series_label": "T 系列",
                    "core_value_proposition": "更稳结果与更低维护负担。",
                    "main_buy_point": "拖后结果可信。",
                    "main_breakpoint": "维护麻烦。",
                    "main_switch_point": "竞品抢主销默认答案。",
                },
            ]
        }
        competition_engineering_artifacts = {
            "competition_engineering_evidence_cards": {
                "series_cards": [
                    {
                        "series_code": "X",
                        "series_label": "X 系列",
                        "value_config_test_cards": [
                            {
                                "value_axis": "免维护闭环",
                                "self_evidence_status": "仅配置可解释，测试待验证",
                                "competitor_switch_evidence_cards": [
                                    {"switch_reading": "当前只能作为潜在能力或待验证信号，不能写成确定性领先。"}
                                ],
                            }
                        ],
                    }
                ]
            }
        }
        future_intelligence_artifacts = {
            "roadmap_stress_test_cards": {
                "cards": [
                    {
                        "value_axis": "清洁结果可信",
                        "roadmap_action": "坚持",
                        "pm_judgment": "坚持体验战，不被参数牵着走。",
                        "do_not_misfollow": "不要把更高温度当主价值。",
                    },
                    {
                        "value_axis": "免维护闭环",
                        "roadmap_action": "加速",
                        "pm_judgment": "自动换滚筒会推动高端托管竞争。",
                        "do_not_misfollow": "不要只做单点自清洁。",
                    },
                ]
            }
        }
        demand_pool_snapshot = {
            "leaf_item_count": 7,
            "theme_rows": [
                {
                    "theme_name": "维护与基站操作",
                    "item_count": 3,
                    "statuses": ["预研中"],
                    "owners": ["产品规划"],
                }
            ],
        }
        cards = insight_bundle.build_report_assembly_cards(
            value_proposition_cards=value_proposition_cards,
            competition_engineering_artifacts=competition_engineering_artifacts,
            future_intelligence_artifacts=future_intelligence_artifacts,
            demand_pool_snapshot=demand_pool_snapshot,
            generated_from="test/integrated_strategy_report.md",
        )
        report = insight_bundle.render_integrated_strategy_report(
            category_name="扫地机",
            time_scope="2026Q2",
            market="中国",
            analysis_goal_text="形成一份产品规划汇报报告",
            report_assembly_cards=cards,
        )
        self.assertIn("一体化产品规划汇报报告", report)
        self.assertIn("## 0. 一页结论", report)
        self.assertIn("## 1. 认清用户现实", report)
        self.assertIn("## 2. 认清产品线任务", report)
        self.assertIn("## 3. 认清竞争压力", report)
        self.assertIn("## 4. 认清未来战局", report)
        self.assertIn("## 5. 需求池排兵布阵与行动指南", report)
        self.assertIn("## 6. 5W2H 行动清单", report)
        self.assertIn("给汇报看的结论", report)
        self.assertIn("给产品规划看的拆解", report)
        self.assertIn("证据锚点与边界", report)
        self.assertIn("value_proposition_cards.json", report)
        self.assertIn("不输出伪 KANO 分数", report)
        self.assertNotIn("75%", report)
        self.assertNotIn("58.2%", report)
        self.assertNotIn("Y=-10", report)


if __name__ == "__main__":
    unittest.main()
