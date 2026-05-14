#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import json
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable


FIELD_ALIASES = {
    "product_version": ["产品版本"],
    "product_subversion": ["产品子版本"],
    "region": ["销售区域"],
    "channel": ["渠道"],
    "msrp": ["建议市场零售价 (CNY)", "建议市场零售价"],
    "channel_units": ["渠道销量"],
    "asp": ["ASP"],
    "net_sales_unit": ["Net Sales"],
    "channel_revenue": ["渠道销售收入"],
    "target_cost_unit": ["目标成本 (含税) (CNY)", "目标成本"],
    "channel_cost": ["渠道销售成本"],
    "channel_gross_margin": ["渠道毛利额"],
}


def load_change_brief(input_dir: Path, output_dir: Path) -> dict[str, Any]:
    candidates = [input_dir / "change_brief.json", output_dir / "change_brief_used.json"]
    for path in candidates:
        if path.exists():
            return json.loads(path.read_text(encoding="utf-8"))
    return {"target_region": "中国", "focus_change_keywords": ["中配"]}


def load_registry(output_dir: Path) -> dict[str, Any]:
    path = output_dir / "source_registry.json"
    if path.exists():
        return json.loads(path.read_text(encoding="utf-8"))
    return {"sources": []}


def find_workbooks(input_dir: Path, output_dir: Path) -> dict[str, Path]:
    registry = load_registry(output_dir)
    result: dict[str, Path] = {}
    for source in registry.get("sources", []):
        role = source.get("role")
        if role in {"before_financial", "after_financial"}:
            result["before" if role == "before_financial" else "after"] = Path(source["path"])
    if "before" in result and "after" in result:
        return result
    for path in input_dir.rglob("*.xlsx"):
        name = path.name.lower()
        if "变更前" in path.name or "before" in name:
            result.setdefault("before", path)
        elif "变更后" in path.name or "after" in name:
            result.setdefault("after", path)
    return result


def header_positions(headers: list[Any]) -> dict[str, int]:
    positions: dict[str, int] = {}
    normalized = {str(value).strip(): index for index, value in enumerate(headers) if value is not None}
    for standard, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            if alias in normalized:
                positions[standard] = normalized[alias]
                break
    return positions


def to_float(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def read_financial_rows(path: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:  # pragma: no cover - environment guard
        raise RuntimeError(f"openpyxl is required to read XLSX: {exc}") from exc

    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    positions = header_positions(list(rows[0]))
    required = ["product_subversion", "region", "channel", "channel_units", "asp", "channel_revenue", "target_cost_unit", "channel_cost", "channel_gross_margin"]
    missing = [field for field in required if field not in positions]
    if missing:
        raise ValueError(f"{path} missing required fields: {', '.join(missing)}")

    parsed: list[dict[str, Any]] = []
    for raw in rows[1:]:
        if not any(value is not None for value in raw):
            continue
        item = {field: raw[index] if index < len(raw) else None for field, index in positions.items()}
        for numeric in ["msrp", "channel_units", "asp", "net_sales_unit", "channel_revenue", "target_cost_unit", "channel_cost", "channel_gross_margin"]:
            item[numeric] = to_float(item.get(numeric))
        parsed.append(item)
    return parsed


def aggregate(rows: Iterable[dict[str, Any]], key_fn) -> dict[str, dict[str, float]]:
    groups: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    for row in rows:
        key = str(key_fn(row))
        units = to_float(row.get("channel_units"))
        groups[key]["rows"] += 1
        groups[key]["units"] += units
        groups[key]["revenue"] += to_float(row.get("channel_revenue"))
        groups[key]["cost"] += to_float(row.get("channel_cost"))
        groups[key]["gross_margin"] += to_float(row.get("channel_gross_margin"))
        groups[key]["asp_x_units"] += to_float(row.get("asp")) * units
        groups[key]["msrp_x_units"] += to_float(row.get("msrp")) * units
        groups[key]["target_cost_x_units"] += to_float(row.get("target_cost_unit")) * units
    return groups


def finalize(group: dict[str, float]) -> dict[str, float]:
    units = group.get("units", 0.0)
    revenue = group.get("revenue", 0.0)
    return {
        "units": units,
        "revenue": revenue,
        "cost": group.get("cost", 0.0),
        "gross_margin": group.get("gross_margin", 0.0),
        "gross_margin_rate": group.get("gross_margin", 0.0) / revenue if revenue else 0.0,
        "weighted_asp": group.get("asp_x_units", 0.0) / units if units else 0.0,
        "weighted_msrp": group.get("msrp_x_units", 0.0) / units if units else 0.0,
        "weighted_target_cost": group.get("target_cost_x_units", 0.0) / units if units else 0.0,
    }


def bridge_rows(before: dict[str, dict[str, float]], after: dict[str, dict[str, float]], group_level: str) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for key in sorted(set(before) | set(after)):
        b = finalize(before.get(key, {}))
        a = finalize(after.get(key, {}))
        rows.append(
            {
                "group_level": group_level,
                "group_key": key,
                "before_units": round(b["units"], 6),
                "after_units": round(a["units"], 6),
                "delta_units": round(a["units"] - b["units"], 6),
                "before_revenue": round(b["revenue"], 6),
                "after_revenue": round(a["revenue"], 6),
                "delta_revenue": round(a["revenue"] - b["revenue"], 6),
                "before_gross_margin": round(b["gross_margin"], 6),
                "after_gross_margin": round(a["gross_margin"], 6),
                "delta_gross_margin": round(a["gross_margin"] - b["gross_margin"], 6),
                "before_gross_margin_rate": round(b["gross_margin_rate"], 6),
                "after_gross_margin_rate": round(a["gross_margin_rate"], 6),
                "delta_gross_margin_rate": round(a["gross_margin_rate"] - b["gross_margin_rate"], 6),
                "before_weighted_asp": round(b["weighted_asp"], 6),
                "after_weighted_asp": round(a["weighted_asp"], 6),
                "before_weighted_msrp": round(b["weighted_msrp"], 6),
                "after_weighted_msrp": round(a["weighted_msrp"], 6),
                "before_weighted_target_cost": round(b["weighted_target_cost"], 6),
                "after_weighted_target_cost": round(a["weighted_target_cost"], 6),
            }
        )
    return rows


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def classify_attribution(row: dict[str, Any], focus_keywords: list[str]) -> str:
    key = str(row["group_key"])
    if any(keyword in key for keyword in focus_keywords):
        return "directly_related"
    if row["group_level"] in {"product_subversion", "product_channel"} and abs(float(row["delta_units"])) > 0:
        return "portfolio_related"
    return "not_attributed"


def build_metric_bridge(input_dir: str | Path, output_dir: str | Path) -> dict[str, Any]:
    input_path = Path(input_dir)
    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)
    brief = load_change_brief(input_path, output_path)
    workbooks = find_workbooks(input_path, output_path)
    if "before" not in workbooks or "after" not in workbooks:
        raise FileNotFoundError("Need both before and after financial workbooks. Use filenames containing 变更前/变更后 or before/after.")

    target_region = brief.get("target_region") or "中国"
    focus_keywords = [str(item) for item in brief.get("focus_change_keywords", ["中配"])]
    before_all = read_financial_rows(workbooks["before"])
    after_all = read_financial_rows(workbooks["after"])
    before = [row for row in before_all if row.get("region") == target_region]
    after = [row for row in after_all if row.get("region") == target_region]

    bridge: list[dict[str, Any]] = []
    bridge.extend(bridge_rows(aggregate(before, lambda row: target_region), aggregate(after, lambda row: target_region), "target_region"))
    bridge.extend(bridge_rows(aggregate(before, lambda row: row.get("product_subversion")), aggregate(after, lambda row: row.get("product_subversion")), "product_subversion"))
    bridge.extend(bridge_rows(aggregate(before, lambda row: row.get("channel")), aggregate(after, lambda row: row.get("channel")), "channel"))
    bridge.extend(bridge_rows(aggregate(before, lambda row: f"{row.get('product_subversion')} / {row.get('channel')}"), aggregate(after, lambda row: f"{row.get('product_subversion')} / {row.get('channel')}"), "product_channel"))
    write_csv(output_path / "metric_bridge.csv", bridge)

    attribution = []
    for row in bridge:
        if row["group_level"] not in {"product_subversion", "product_channel", "channel"}:
            continue
        tag = classify_attribution(row, focus_keywords)
        attribution.append(
            {
                "group_level": row["group_level"],
                "group_key": row["group_key"],
                "delta_units": row["delta_units"],
                "delta_revenue": row["delta_revenue"],
                "delta_gross_margin": row["delta_gross_margin"],
                "attribution_tag": tag,
                "note": "Matches focus change keyword" if tag == "directly_related" else "Portfolio movement; validate causality before claiming attribution",
            }
        )
    write_csv(output_path / "attribution_table.csv", attribution)

    total_before = finalize(aggregate(before, lambda row: "total").get("total", {}))
    total_after = finalize(aggregate(after, lambda row: "total").get("total", {}))
    new_rows = [row for row in after if any(keyword in str(row.get("product_subversion")) for keyword in focus_keywords)]
    new_summary = finalize(aggregate(new_rows, lambda row: "new_sku").get("new_sku", {}))
    summary = {
        "target_region": target_region,
        "focus_change_keywords": focus_keywords,
        "before_workbook": str(workbooks["before"]),
        "after_workbook": str(workbooks["after"]),
        "total": {"before": total_before, "after": total_after, "delta": {key: total_after[key] - total_before[key] for key in total_before}},
        "new_sku": new_summary,
        "new_sku_contribution": {
            "units": new_summary["units"] / (total_after["units"] - total_before["units"]) if total_after["units"] != total_before["units"] else None,
            "revenue": new_summary["revenue"] / (total_after["revenue"] - total_before["revenue"]) if total_after["revenue"] != total_before["revenue"] else None,
            "gross_margin": new_summary["gross_margin"] / (total_after["gross_margin"] - total_before["gross_margin"]) if total_after["gross_margin"] != total_before["gross_margin"] else None,
        },
    }
    (output_path / "metric_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary


def main() -> None:
    parser = argparse.ArgumentParser(description="Build before/after metric bridge and attribution table.")
    parser.add_argument("--input-dir", required=True)
    parser.add_argument("--output-dir", required=True)
    args = parser.parse_args()
    summary = build_metric_bridge(args.input_dir, args.output_dir)
    print(json.dumps({"target_region": summary["target_region"], "output_dir": args.output_dir}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
