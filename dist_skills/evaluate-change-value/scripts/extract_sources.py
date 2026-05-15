#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from typing import Any


SUPPORTED_SUFFIXES = {".docx", ".xlsx", ".xlsm", ".csv", ".md", ".txt", ".json", ".png", ".jpg", ".jpeg"}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def infer_role(path: Path) -> str:
    name = path.name.lower()
    if path.suffix.lower() in {".xlsx", ".xlsm", ".csv"}:
        if "变更前" in path.name or "before" in name:
            return "before_financial"
        if "变更后" in path.name or "after" in name:
            return "after_financial"
        return "financial_unknown"
    if path.suffix.lower() == ".docx":
        if "用户" in path.name or "user" in name or "voc" in name:
            return "user_research"
        return "document"
    if path.suffix.lower() in {".png", ".jpg", ".jpeg"}:
        return "config_or_evidence_image"
    if path.name == "change_brief.json":
        return "change_brief"
    return "supporting_text"


def extract_docx(path: Path) -> dict[str, Any]:
    try:
        from docx import Document
    except Exception as exc:  # pragma: no cover - environment guard
        return {"error": f"python-docx is required to read DOCX: {exc}"}

    doc = Document(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    keywords = ["用户", "画像", "家庭", "宠", "房屋", "购买", "痛点", "维护", "续航", "清洁", "竞品"]
    evidence = []
    for index, text in enumerate(paragraphs, start=1):
        if any(keyword in text for keyword in keywords):
            evidence.append({"paragraph": index, "text": text[:500]})
        if len(evidence) >= 80:
            break
    return {
        "paragraph_count": len(paragraphs),
        "table_count": len(doc.tables),
        "evidence": evidence,
    }


def extract_xlsx(path: Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:  # pragma: no cover - environment guard
        return {"error": f"openpyxl is required to read XLSX: {exc}"}

    workbook = load_workbook(path, read_only=False, data_only=True)
    sheets = []
    for sheet in workbook.worksheets:
        header = [sheet.cell(1, col).value for col in range(1, min(sheet.max_column, 80) + 1)]
        sheets.append(
            {
                "title": sheet.title,
                "rows": sheet.max_row,
                "columns": sheet.max_column,
                "header": [value for value in header if value is not None],
            }
        )
    return {"sheets": sheets}


def extract_text(path: Path) -> dict[str, Any]:
    try:
        text = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        text = path.read_text(encoding="utf-8", errors="replace")
    return {"chars": len(text), "preview": text[:1000]}


def scan_sources(input_dir: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    sources: list[dict[str, Any]] = []
    user_evidence: list[dict[str, Any]] = []
    for path in sorted(input_dir.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED_SUFFIXES:
            continue
        item: dict[str, Any] = {
            "path": str(path.resolve()),
            "name": path.name,
            "suffix": path.suffix.lower(),
            "role": infer_role(path),
            "bytes": path.stat().st_size,
            "sha256": sha256_file(path),
        }
        if path.suffix.lower() == ".docx":
            docx = extract_docx(path)
            item["docx"] = {k: v for k, v in docx.items() if k != "evidence"}
            for evidence in docx.get("evidence", []):
                user_evidence.append({"source": path.name, **evidence})
        elif path.suffix.lower() in {".xlsx", ".xlsm"}:
            item["workbook"] = extract_xlsx(path)
        elif path.suffix.lower() in {".md", ".txt", ".json"}:
            item["text"] = extract_text(path)
        sources.append(item)
    return sources, user_evidence


def extract_sources(input_dir: str | Path, output_dir: str | Path) -> dict[str, Any]:
    input_path = Path(input_dir)
    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)
    sources, user_evidence = scan_sources(input_path)
    registry = {"input_dir": str(input_path.resolve()), "source_count": len(sources), "sources": sources}
    (output_path / "source_registry.json").write_text(
        json.dumps(registry, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (output_path / "user_evidence.json").write_text(
        json.dumps({"evidence_count": len(user_evidence), "evidence": user_evidence}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return registry


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract source registry and user evidence from a change-evaluation input folder.")
    parser.add_argument("--input-dir", required=True)
    parser.add_argument("--output-dir", required=True)
    args = parser.parse_args()
    registry = extract_sources(args.input_dir, args.output_dir)
    print(json.dumps({"source_count": registry["source_count"], "output_dir": args.output_dir}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
